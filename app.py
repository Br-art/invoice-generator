import os
import io
import json
import re
from datetime import datetime, date
from typing import List, Tuple, Optional

import streamlit as st
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from fpdf import FPDF

# =========================
# Fixed cell mapping (template)
# =========================
DELIVER_TO_CELL = "G19"
DATE_CELL = "I6"
DOCNO_CELL = "I12"

# Single-ticket cells
ROUTE_CELL = "E35"
PRICE_CELL = "I35"

# Flights (dates) table area
FLIGHT_START_ROW = 42
FLIGHT_MAX_ROWS = 40

# Passenger names list
PASSENGER_LABEL_CELL = "C45"
PASSENGER_START_ROW = 47
PASSENGER_MAX_ROWS = 9

# Subtotal / Total
SUBTOTAL_LABEL_CELL, SUBTOTAL_CELL = "H56", "I56"
TOTAL_LABEL_CELL, TOTAL_CELL = "H59", "I59"

# Multi-ticket summary
SUMMARY_START_ROW = 36
SUMMARY_MAX_ROWS = 10

# Currency symbols
CURRENCY_SYMBOL = {
    "ZMW": "K",
    "ZAR": "R",
    "USD": "$",
    "GBP": "£",
    "EUR": "€",
}

# Persistence
SETTINGS_FILE = "settings.json"
TEMPLATE_FILE = "template.xlsx"

# =========================
# Settings helpers
# =========================
def load_settings() -> dict:
    if os.path.exists(SETTINGS_FILE):
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}

def save_settings(d: dict) -> None:
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(d, f)

settings = load_settings()

# =========================
# Parsing helpers
# =========================
def parse_reference(txt: str) -> Optional[str]:
    pattern = re.compile(r"\b(?=[A-Z0-9]{5,6}\b)(?=.*[A-Z])[A-Z0-9]{5,6}\b")
    lines = [ln.strip() for ln in txt.splitlines() if ln.strip()]
    tail = "\n".join(lines[-12:]) if lines else txt

    m_last = None
    for m in pattern.finditer(tail):
        m_last = m
    if m_last:
        return m_last.group(0)

    m2_last = None
    for m in pattern.finditer(txt):
        m2_last = m
    return m2_last.group(0) if m2_last else None

def parse_total_currency(txt: str) -> Tuple[Optional[str], Optional[float]]:
    m = re.search(r"Total\s+([A-Z]{3})\s+([0-9]+(?:\.[0-9]{1,2})?)", txt)
    if not m:
        return None, None
    return m.group(1), float(m.group(2))

def parse_flights(txt: str) -> List[Tuple[str, str]]:
    flights: List[Tuple[str, str]] = []
    for line in txt.splitlines():
        m = re.search(r"(\d{2}[A-Z]{3}\d{2})\s*([A-Z]{3})\s+([A-Z]{3})", line)
        if m:
            ddmmmyy, orig, dest = m.groups()
            try:
                d = datetime.strptime(ddmmmyy, "%d%b%y").date()
                date_str = d.strftime("%d.%m.%Y")
            except Exception:
                date_str = ddmmmyy
            flights.append((f"{orig}-{dest}", date_str))
    return flights

def pretty_name(slash_name: str) -> str:
    parts = slash_name.split("/")
    if len(parts) != 2:
        return slash_name.title()

    last = parts[0].title()
    given_raw = parts[1]
    title_map = {
        "MR": "Mr",
        "MS": "Ms",
        "MRS": "Mrs",
        "MSTR": "Mr",
        "MISS": "Miss",
    }
    is_inf = ".IN" in slash_name
    m = re.match(r"^([A-Z]+?)(MR|MS|MRS|MSTR|MISS)?(?:\.IN\d+)?$", given_raw)
    title = "Mr/Ms"
    given = given_raw.title()

    if m:
        core, t = m.groups()
        given = core.title()
        if t and t in title_map:
            title = title_map[t]

    return f"{title} {given} {last}" + (" (INF)" if is_inf else "")

def parse_passengers(txt: str) -> List[str]:
    pax: List[str] = []
    for m in re.finditer(r"TICKET ISSUED .*?\s([A-Z]+/[A-Z]+[A-Z\.0-9]*)", txt):
        pax.append(pretty_name(m.group(1)))

    if pax:
        out, seen = [], set()
        for p in pax:
            if p not in seen:
                seen.add(p)
                out.append(p)
        return out

    for m in re.finditer(r"\d+\.\d+([A-Z]+)/([A-Z]+[A-Z\.0-9]*)", txt):
        pax.append(pretty_name(m.group(1) + "/" + m.group(2)))

    out, seen = [], set()
    for p in pax:
        if p not in seen:
            seen.add(p)
            out.append(p)
    return out

def split_tickets(big_text: str) -> List[str]:
    s = big_text.strip()
    if not s:
        return []
    parts = re.split(r"\n\s*---\s*\n", s)
    if len(parts) > 1:
        return [p.strip() for p in parts if p.strip()]
    parts = re.split(r"(?:\r?\n){2,}", s)
    return [p.strip() for p in parts if p.strip()]

def parse_one_ticket(txt: str) -> dict:
    ref = parse_reference(txt) or "REFXXX"
    ccy, total = parse_total_currency(txt)
    flights = parse_flights(txt)
    pax = parse_passengers(txt)
    return {
        "ref": ref,
        "currency": ccy or "ZMW",
        "total": float(total or 0.0),
        "flights": flights,
        "pax": pax,
    }

def ticket_route_string(flights: List[Tuple[str, str]]) -> str:
    if not flights:
        return "N/A"
    first = flights[0][0].split("-")
    chain = [first[0], first[1]] if len(first) == 2 else [flights[0][0]]
    for seg, _ in flights[1:]:
        parts = seg.split("-")
        chain.append(parts[1] if len(parts) == 2 else seg)
    return "-".join(chain)

# =========================
# Excel writers
# =========================
def clear_block(ws, start_row: int, rows: int, cols: List[str]) -> None:
    for r in range(start_row, start_row + rows):
        for c in cols:
            ws[f"{c}{r}"] = ""

def write_flight_dates(ws, ticket_items: List[dict]) -> None:
    for r in range(FLIGHT_START_ROW, FLIGHT_START_ROW + FLIGHT_MAX_ROWS):
        ws[f"C{r}"], ws[f"E{r}"] = "", ""
    row = FLIGHT_START_ROW
    for t in ticket_items:
        for route, d in (t["flights"] or [("N/A", "")]):
            if row >= FLIGHT_START_ROW + FLIGHT_MAX_ROWS:
                ws[f"C{FLIGHT_START_ROW + FLIGHT_MAX_ROWS - 1}"] = "(+more)"
                return
            ws[f"C{row}"] = route
            ws[f"E{row}"] = d
            row += 1

def write_passengers(ws, pax: List[str]) -> None:
    ws[PASSENGER_LABEL_CELL] = "Passengers:"
    clear_block(ws, PASSENGER_START_ROW, PASSENGER_MAX_ROWS, ["C"])
    for i, name in enumerate(pax[:PASSENGER_MAX_ROWS]):
        ws[f"C{PASSENGER_START_ROW + i}"] = name

def write_ticket_summary(ws, ticket_items: List[dict]) -> None:
    for r in range(SUMMARY_START_ROW, SUMMARY_START_ROW + SUMMARY_MAX_ROWS):
        for c in ["E", "F", "I"]:
            ws[f"{c}{r}"] = ""
    row = SUMMARY_START_ROW
    for t in ticket_items[:SUMMARY_MAX_ROWS]:
        sym = CURRENCY_SYMBOL.get(t["currency"], "K")
        ws[f"E{row}"] = ticket_route_string(t["flights"])
        ws[f"F{row}"] = t["ref"]
        ws[f"I{row}"] = f"{sym}{t['total']:,.2f}"
        row += 1

def configure_print_settings(ws) -> None:
    ws.print_area = "A1:I75"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

def apply_to_workbook(wb: Workbook, ticket_text: str, customer: str) -> Workbook:
    ws = wb.active
    ref = parse_reference(ticket_text) or "REFXXX"
    ccy, total = parse_total_currency(ticket_text)
    flights = parse_flights(ticket_text)
    pax = parse_passengers(ticket_text)

    sym = CURRENCY_SYMBOL.get(ccy or "ZMW", "K")
    amount_str = f"{sym}{(total or 0.0):,.2f}"

    ws[DELIVER_TO_CELL] = customer.strip()
    ws[DATE_CELL] = date.today()
    ws[DOCNO_CELL] = f"{ref}{date.today():%d%m%y}"

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "ABBNWU" in cell.value:
                cell.value = cell.value.replace("ABBNWU", ref)

    write_flight_dates(ws, [{"flights": flights}])
    if flights:
        unique_routes = []
        for r, _ in flights:
            if r not in unique_routes:
                unique_routes.append(r)
        ws[ROUTE_CELL] = " – ".join(unique_routes)
    else:
        ws[ROUTE_CELL] = ""

    ws[PRICE_CELL] = amount_str
    ws[SUBTOTAL_LABEL_CELL] = "Sub Total"
    ws[TOTAL_LABEL_CELL] = "Total Payable"
    ws[SUBTOTAL_CELL] = amount_str
    ws[TOTAL_CELL] = amount_str

    write_passengers(ws, pax)
    configure_print_settings(ws)
    return wb

def apply_to_workbook_multi(wb: Workbook, big_ticket_text: str, customer: str) -> Workbook:
    ws = wb.active
    chunks = split_tickets(big_ticket_text)
    items = [parse_one_ticket(c) for c in chunks if c]

    first_ref = items[0]["ref"] if items else "REFXXX"
    ws[DOCNO_CELL] = f"{first_ref}{date.today():%d%m%y}"
    ws[DELIVER_TO_CELL] = customer.strip()
    ws[DATE_CELL] = date.today()

    ref_list = ", ".join([x["ref"] for x in items])
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "ABBNWU" in cell.value:
                cell.value = cell.value.replace("ABBNWU", ref_list)

    ws["E35"], ws["I35"] = "", ""
    write_ticket_summary(ws, items)
    write_flight_dates(ws, items)

    union = []
    for t in items:
        for p in t["pax"]:
            if p not in union:
                union.append(p)
    write_passengers(ws, union)

    grand = sum(t["total"] for t in items)
    first_ccy = items[0]["currency"] if items else "ZMW"
    sym = CURRENCY_SYMBOL.get(first_ccy, "K")
    amt = f"{sym}{grand:,.2f}"

    ws[SUBTOTAL_LABEL_CELL] = "Sub Total"
    ws[TOTAL_LABEL_CELL] = "Total Payable"
    ws[SUBTOTAL_CELL] = amt
    ws[TOTAL_CELL] = amt
    configure_print_settings(ws)
    return wb

# =========================
# Pure Python PDF Generator
# =========================
def create_pure_pdf(customer: str, items: List[dict]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    
    # Title
    pdf.set_font("helvetica", "B", 20)
    pdf.cell(0, 10, "INVOICE", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)
    
    # Customer and Date Details
    pdf.set_font("helvetica", "", 12)
    pdf.cell(100, 8, f"Deliver To: {customer}")
    pdf.cell(90, 8, f"Date: {date.today().strftime('%d %b %Y')}", align="R", new_x="LMARGIN", new_y="NEXT")
    
    first_ref = items[0]["ref"] if items else "REFXXX"
    doc_no = f"{first_ref}{date.today():%d%m%y}"
    pdf.cell(100, 8, f"Invoice No: {doc_no}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)
    
    # Table Header
    pdf.set_font("helvetica", "B", 12)
    pdf.cell(40, 10, "Reference", border="B")
    pdf.cell(100, 10, "Route", border="B")
    pdf.cell(50, 10, "Amount", border="B", align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    
    # Table Rows
    grand_total = 0
    first_ccy = items[0]["currency"] if items else "ZMW"
    sym = CURRENCY_SYMBOL.get(first_ccy, "K")
    
    pdf.set_font("helvetica", "", 11)
    for t in items:
        route_str = ticket_route_string(t["flights"])
        amt_str = f"{CURRENCY_SYMBOL.get(t['currency'], 'K')}{t['total']:,.2f}"
        grand_total += t["total"]
        
        pdf.cell(40, 10, t["ref"])
        pdf.cell(100, 10, route_str)
        pdf.cell(50, 10, amt_str, align="R", new_x="LMARGIN", new_y="NEXT")
        
        # Add passengers under each route
        if t["pax"]:
            pdf.set_font("helvetica", "I", 10)
            pax_str = "Passengers: " + ", ".join(t["pax"])
            pdf.cell(40, 6, "") # spacer
            pdf.multi_cell(100, 6, pax_str, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("helvetica", "", 11)
            pdf.ln(2)
            
    # Totals
    pdf.ln(5)
    pdf.set_font("helvetica", "B", 14)
    total_str = f"{sym}{grand_total:,.2f}"
    pdf.cell(140, 10, "Total Payable:", align="R")
    pdf.cell(50, 10, total_str, align="R", new_x="LMARGIN", new_y="NEXT")
    
    return bytes(pdf.output())

# =========================
# Streamlit UI
# =========================
st.set_page_config(page_title="Invoice Generator", page_icon="📄")
st.title("Invoice Generator")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(BASE_DIR, TEMPLATE_FILE)

if not os.path.exists(TEMPLATE_PATH):
    st.error("template.xlsx is missing from the app folder.")
    st.stop()

if "deliver_to" not in st.session_state:
    st.session_state.deliver_to = settings.get("last_deliver_to", "")

customer = st.text_input("Deliver To (customer/company)", value=st.session_state.deliver_to)
st.session_state.deliver_to = customer

ticket_text = st.text_area("Paste ticket text", height=280)

if st.button("Generate Invoice", type="primary"):
    if not customer.strip() or not ticket_text.strip():
        st.error("Please fill in both the customer name and ticket text.")
        st.stop()

    # Parse data for both Excel and PDF
    chunks = split_tickets(ticket_text)
    items = [parse_one_ticket(c) for c in chunks if c] if len(chunks) >= 2 else [parse_one_ticket(ticket_text)]
    
    # 1. Generate Excel
    wb = load_workbook(TEMPLATE_PATH)
    if len(chunks) >= 2:
        wb = apply_to_workbook_multi(wb, ticket_text, customer)
    else:
        wb = apply_to_workbook(wb, ticket_text, customer)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_bytes = excel_buffer.getvalue()

    # 2. Generate pure Python PDF
    pdf_bytes = create_pure_pdf(customer, items)

    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", customer.strip())
    base_name = f"Invoice_{safe}_{date.today():%Y%m%d}"

    st.success("Invoice generated successfully! ✅")

    col1, col2 = st.columns(2)
    with col1:
        st.download_button(
            "⬇️ Download Excel Invoice",
            data=excel_bytes,
            file_name=f"{base_name}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    with col2:
        st.download_button(
            "⬇️ Download PDF Invoice",
            data=pdf_bytes,
            file_name=f"{base_name}.pdf",
            mime="application/pdf",
        )

    save_settings({"last_deliver_to": customer})
